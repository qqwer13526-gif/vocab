"""从 tests/fixtures/vocab-sample.tsv 生成同名 .xlsx，用来测"应用直接读 Excel"。

刻意做成和用户真表一样的形态：第 1 行标题（合并 A1:F1）、第 2 行统计、第 3 行空、
第 4 行表头、第 5 行起数据；纯数字的单元格写成数字（这样能覆盖 xlsx 里"数字不查共享字符串表"的分支）。

用法：python tool/make_sample_xlsx.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "tests" / "fixtures" / "vocab-sample.tsv"
OUT = ROOT / "tests" / "fixtures" / "vocab-sample.xlsx"


def main() -> int:
    rows = [line.split("\t") for line in SRC.read_text(encoding="utf-8").splitlines()]
    wb = Workbook()
    ws = wb.active
    ws.title = "词汇表"
    for r, cells in enumerate(rows, start=1):
        for c, value in enumerate(cells, start=1):
            text = value.strip()
            if not text:
                continue
            ws.cell(row=r, column=c, value=int(text) if text.isdigit() else text)
    ws.merge_cells("A1:F1")
    wb.save(OUT)
    print(f"已生成 {OUT.relative_to(ROOT)}（{OUT.stat().st_size} 字节，{len(rows)} 行，工作表「{ws.title}」）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
